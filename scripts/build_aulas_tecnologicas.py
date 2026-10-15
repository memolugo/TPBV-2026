"""
Genera dashboard/aulas_tecnologicas.js a partir de la hoja "Concentrado" de
'Proyecto Escuelas entrega Computadoras.xlsx' (aulas de computo entregadas por
ATDEM en escuelas de los poligonos TPBV, ver boletin 06206 del 30/04/2026).

  - La ubicacion de cada escuela viene como link corto de Google Maps
    (maps.app.goo.gl); se resuelve siguiendo la redireccion y leyendo
    !3d<lat>!4d<lon> de la URL final. Los resultados se guardan en
    scripts/aulas_tecnologicas_coords.json para que reconstruir no dependa de
    la red (solo se consultan los links que no esten en ese cache).
  - El poligono TPBV de cada escuela se calcula con punto-en-poligono contra
    TPBV 2026.geojson (Propuesta Actual).
  - NO se exportan nombres de coordinadores/promotores ni telefonos: el
    tablero se publica en GitHub Pages.
"""
import json
import math
import re
import time
import unicodedata
import urllib.request
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
XLSX = BASE / "Proyecto Escuelas entrega Computadoras.xlsx"
GEOJSON = BASE / "TPBV 2026.geojson"
NOMBRES_POL = BASE / "dashboard" / "poligonos_actual_nombres.js"
CACHE = Path(__file__).resolve().parent / "aulas_tecnologicas_coords.json"
OUT = BASE / "dashboard" / "aulas_tecnologicas.js"

# Columnas de la hoja "Concentrado" (fila 3 = encabezados, datos desde la 4)
COL = {"n": 0, "desc": 1, "region": 2, "computadoras": 4, "url": 9,
       "fase": 11, "fechaIebem": 15, "fechaAtdem": 17}

CCT_RE = re.compile(r"\b17[A-Z]{3}\d{3,4}[A-Z]\b")

# Tipo de servicio segun las letras 3-5 de la CCT (catalogo SEP)
NIVEL_POR_CCT = {
    "DPR": "Primaria", "DPB": "Primaria indígena", "KPR": "Primaria comunitaria (CONAFE)",
    "DES": "Secundaria general", "DST": "Secundaria técnica", "DTV": "Telesecundaria",
    "EMS": "Media superior (EMSAD)",
}

FASE_NORMALIZADA = {
    "primera fase": "Primera Fase", "segunda fase": "Segunda Fase",
    "conflicto de turno": "Conflicto de turno",
}

# Filas cuyo texto "Municipio, Localidad, Escuela, CCT" no trae las comas
# esperadas en la hoja; se corrigen aqui en vez de adivinar.
CORRECCIONES_DESC = {
    33: ("Zacatepec", "Plutarco Elías Calles", "16 de Septiembre"),
}

# Escuelas cuyo link de Google Maps en la hoja apunta a otra escuela (p. ej.
# un homonimo). Se usa la coordenada de los datos publicos SEP por CCT.
COORDENADAS_CORREGIDAS = {
    # Link de la hoja -> "Escuela Primaria Marti De Chinameca" cerca de Cuautla
    # (18.7178, -98.8998); la CCT esta en Lazaro Cardenas 45, Chinameca, Ayala
    # (SEP 2018/2023 via escuelasmex.com/directorio/17DPR0486P).
    "17DPR0486P": (18.622483, -98.995379),
}


def norm(s):
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower().strip()


def limpiar(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def parse_desc(n, desc):
    cct_match = CCT_RE.search(desc.upper())
    cct = cct_match.group(0) if cct_match else ""
    if n in CORRECCIONES_DESC:
        municipio, localidad, escuela = CORRECCIONES_DESC[n]
        return municipio, localidad, escuela, cct
    sin_cct = CCT_RE.sub("", desc.upper()) if cct else desc
    # Recupera el texto original (con minusculas) quitando la CCT por posicion
    if cct_match:
        idx = desc.upper().find(cct)
        sin_cct = desc[:idx] + desc[idx + len(cct):]
    partes = [limpiar(p) for p in sin_cct.split(",") if limpiar(p)]
    municipio = partes[0] if partes else ""
    localidad = partes[1] if len(partes) > 1 else ""
    escuela = ", ".join(partes[2:]) if len(partes) > 2 else ""
    return municipio, localidad, escuela, cct


def resolver_coordenada(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    final = urllib.request.urlopen(req, timeout=30).geturl()
    # Formatos vistos: lugar (!3d<lat>!4d<lon>), busqueda por coordenada
    # (/maps/search/<lat>,+<lon>) y ruta (/maps/dir/<origen>/<destino>/, se
    # toma el destino = ultimo par).
    lugar = re.search(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)", final)
    busqueda = re.search(r"/maps/search/(-?\d+\.\d+),\+?(-?\d+\.\d+)", final)
    ruta = re.findall(r"/(-?\d+\.\d+),\+?(-?\d+\.\d+)(?=/)", final) if "/maps/dir/" in final else []
    vista = re.search(r"@(-?\d+\.\d+),(-?\d+\.\d+)", final)
    par = (lugar.groups() if lugar else busqueda.groups() if busqueda
           else ruta[-1] if ruta else vista.groups() if vista else None)
    if not par:
        raise ValueError(f"sin coordenada en la URL final: {final[:160]}")
    return round(float(par[0]), 6), round(float(par[1]), 6)


def cargar_poligonos():
    gj = json.loads(GEOJSON.read_text(encoding="utf-8"))
    src = NOMBRES_POL.read_text(encoding="utf-8")
    nombres = json.loads(re.search(r"= (\{.*?\});", src, re.S).group(1))
    poligonos = []
    for f in gj["features"]:
        geom = f["geometry"]
        polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
        pol = int(f["properties"]["Pol"])  # en el geojson viene como texto
        poligonos.append((pol, nombres.get(str(pol), f["properties"]["Nom_pol"]), polys))
    return poligonos


def dentro_de_anillo(ring, lon, lat):
    dentro = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi, yi = ring[i][:2]
        xj, yj = ring[j][:2]
        if (yi > lat) != (yj > lat) and lon < (xj - xi) * (lat - yi) / (yj - yi) + xi:
            dentro = not dentro
        j = i
    return dentro


def distancia_km_a_segmento(lon, lat, a, b):
    # Aproximacion plana (suficiente a escala de Morelos): grados -> km con
    # el coseno de la latitud para la longitud.
    kx, ky = 111.32 * math.cos(math.radians(lat)), 110.57
    px, py = lon * kx, lat * ky
    ax, ay, bx, by = a[0] * kx, a[1] * ky, b[0] * kx, b[1] * ky
    dx, dy = bx - ax, by - ay
    t = 0 if dx == dy == 0 else max(0, min(1, ((px - ax) * dx + (py - ay) * dy) / (dx * dx + dy * dy)))
    return math.hypot(px - (ax + t * dx), py - (ay + t * dy))


def poligono_de(poligonos, lat, lon):
    """(poligono, nombre, dentro, km_al_borde): el que contiene el punto o,
    si ninguno lo contiene, el mas cercano y su distancia al borde."""
    mejor = (None, None, False, float("inf"))
    for pol, nombre, polys in poligonos:
        for poly in polys:
            outer, *holes = poly
            if dentro_de_anillo(outer, lon, lat) and not any(dentro_de_anillo(h, lon, lat) for h in holes):
                return pol, nombre, True, 0.0
            km = min(distancia_km_a_segmento(lon, lat, outer[i], outer[i + 1]) for i in range(len(outer) - 1))
            if km < mejor[3]:
                mejor = (pol, nombre, False, km)
    return mejor[0], mejor[1], False, round(mejor[3], 2)


def leer_mobiliario(ws):
    """Totales de mobiliario adquirido (celda P1 de Concentrado, p. ej.
    "Total\\n372 Sillas\\n294 Mesas\\n664 Tarjetas Wifi")."""
    texto = str(ws["P1"].value or "")

    def total(patron):
        m = re.search(r"(\d+)\s*" + patron, texto, re.I)
        return int(m.group(1)) if m else None

    return {"sillas": total("sillas"), "mesas": total("mesas"), "tarjetasWifi": total("tarjetas")}


def main():
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    poligonos = cargar_poligonos()
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Concentrado"]

    aulas, errores = [], []
    for row in ws.iter_rows(min_row=4, values_only=True):
        n = row[COL["n"]]
        if not isinstance(n, int) or not row[COL["desc"]]:
            continue
        municipio, localidad, escuela, cct = parse_desc(n, limpiar(row[COL["desc"]]))
        url = limpiar(row[COL["url"]])

        if url and url not in cache:
            try:
                cache[url] = resolver_coordenada(url)
                time.sleep(0.5)
            except Exception as e:  # red o link roto: se reporta y la escuela queda sin punto
                errores.append(f"#{n} {escuela}: {e}")
        lat, lon = COORDENADAS_CORREGIDAS.get(cct) or cache.get(url, (None, None))
        pol, nombre_pol, dentro, km = (poligono_de(poligonos, lat, lon) if lat is not None
                                       else (None, None, False, None))

        fase_txt = norm(row[COL["fase"]] or "")
        aulas.append({
            "n": n,
            "municipio": municipio,
            "localidad": localidad,
            "escuela": escuela,
            "cct": cct,
            "nivel": NIVEL_POR_CCT.get(cct[2:5], "Otro") if cct else "Sin CCT",
            "region": limpiar(row[COL["region"]]),
            "computadoras": int(row[COL["computadoras"]] or 0),
            "fase": FASE_NORMALIZADA.get(fase_txt, limpiar(row[COL["fase"]])),
            "fechaEntregaIebem": limpiar(row[COL["fechaIebem"]]),
            "fechaInstalacion": limpiar(row[COL["fechaAtdem"]]),
            "lat": lat,
            "lon": lon,
            "mapsUrl": url,
            # dentroPoligono=False -> poligono/nombrePoligono son el TPBV mas
            # cercano y kmPoligono la distancia a su borde.
            "poligono": pol,
            "nombrePoligono": nombre_pol,
            "dentroPoligono": dentro,
            "kmPoligono": km,
        })

    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    OUT.write_text(
        "// Generado automaticamente por scripts/build_aulas_tecnologicas.py\n"
        "// Fuente: 'Proyecto Escuelas entrega Computadoras.xlsx', hoja Concentrado\n"
        f"const TPBV_AULAS_TECNOLOGICAS = {json.dumps(aulas, ensure_ascii=False)};\n"
        f"const TPBV_AULAS_MOBILIARIO = {json.dumps(leer_mobiliario(ws), ensure_ascii=False)};\n",
        encoding="utf-8")

    sin_punto = [a for a in aulas if a["lat"] is None]
    fuera = [a for a in aulas if a["lat"] is not None and not a["dentroPoligono"]]
    print(f"Aulas: {len(aulas)} | computadoras: {sum(a['computadoras'] for a in aulas)}")
    print(f"Sin coordenada: {len(sin_punto)} | fuera de poligono TPBV: {len(fuera)}")
    for a in fuera:
        print(f"  fuera: #{a['n']} {a['municipio']}, {a['localidad']} - {a['escuela']}"
              f" -> {a['kmPoligono']} km del poligono {a['poligono']} {a['nombrePoligono']}")
    for e in errores:
        print("  ERROR", e)


if __name__ == "__main__":
    main()
