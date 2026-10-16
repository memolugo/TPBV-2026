"""Convierte la plantilla traducida en el diccionario que usa el tablero.

Entrada:  traducciones/Traduccion_nahuatl_tablero.xlsx (hoja "Textos")
Salida:   dashboard/traducciones_na.js  (const TPBV_TRADUCCIONES_NA)

Valida que cada traducción conserve los mismos marcadores {n} y {d} que el
texto en español; las filas que no cumplen se omiten y se reportan.

Uso:  python scripts/build_traducciones_na.py
Luego: python scripts/embed_dashboard.py
"""
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "traducciones" / "Traduccion_nahuatl_tablero.xlsx"
SALIDA = RAIZ / "dashboard" / "traducciones_na.js"
HOJA = "Textos"
COL_CLAVE, COL_TRAD = 4, 5

RE_MARCA = re.compile(r"\{[nd]\}")


def normalizar(texto: str) -> str:
    return re.sub(r"\s+", " ", str(texto)).strip()


def marcas(texto: str) -> Counter:
    return Counter(RE_MARCA.findall(texto))


def leer_traducciones() -> tuple[dict, list]:
    ws = load_workbook(ENTRADA, read_only=True)[HOJA]
    dicc, errores = {}, []
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        clave, trad = fila[COL_CLAVE - 1], fila[COL_TRAD - 1]
        if not clave or not trad or not str(trad).strip():
            continue
        clave, trad = normalizar(clave), normalizar(trad)
        if marcas(clave) != marcas(trad):
            errores.append(f"fila {n}: marcadores distintos -> «{clave[:60]}» / «{trad[:60]}»")
            continue
        dicc[clave] = trad
    return dicc, errores


def main() -> None:
    if not ENTRADA.exists():
        raise SystemExit(f"No existe {ENTRADA}. Ejecute antes scripts/build_plantilla_nahuatl.py")
    dicc, errores = leer_traducciones()
    cuerpo = json.dumps(dict(sorted(dicc.items())), ensure_ascii=False, indent=1)
    SALIDA.write_text(
        "// Diccionario espanol -> nahuatl del tablero. Generado por scripts/build_traducciones_na.py\n"
        "// desde traducciones/Traduccion_nahuatl_tablero.xlsx: no editar a mano.\n"
        f"const TPBV_TRADUCCIONES_NA = {cuerpo};\n",
        encoding="utf-8",
    )
    print(f"{SALIDA.relative_to(RAIZ)}: {len(dicc)} traducciones")
    for e in errores:
        print("  OMITIDA", e)


if __name__ == "__main__":
    main()
