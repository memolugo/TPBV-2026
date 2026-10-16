"""Genera la plantilla de traducción al náhuatl para el traductor.

Entrada:  traducciones/textos_extraidos.json  (scripts/extraer_textos_nahuatl.js)
Salida:   traducciones/Traduccion_nahuatl_tablero.xlsx

Si la plantilla ya existe, conserva las traducciones capturadas (por texto en
español), así se puede regenerar cuando cambien los textos del tablero sin
perder el trabajo del traductor. Se prellenan los términos que ya tienen
traducción oficial en "Traducción transformación digital.docx" (portal
morelos.gob.mx/nahuatl).

Uso:  python scripts/build_plantilla_nahuatl.py
Luego, con la plantilla traducida: python scripts/build_traducciones_na.py
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "traducciones" / "textos_extraidos.json"
SALIDA = RAIZ / "traducciones" / "Traduccion_nahuatl_tablero.xlsx"
HOJA = "Textos"

# Traducciones oficiales del portal (docx "Traducción transformación digital").
GLOSARIO = {
    "Inicio": "Pehua",
    "Gobierno": "Kaltlayekana",
    "La tierra que nos une": "Nin tlale tech sepaneloa",
    "Prensa": "Tlatolnonotza",
    "Trámites y servicios": "Kalakis amamej wan tleonon tlakapalewisque",
    "Atención ciudadana": "Tikinkaquis nin sepanchante",
    "Buen vivir": "Kuale tinemisque",
    "Humanismo mexicano": "Tikinequis kenime icej mexicano",
    "Ética y fraternidad": "Tekitis ka moyolo wan tenequis",
}
# Vocabulario de referencia de la variante de Hueyapan (Tetela del Volcán):
# Montero Baeza, "Ejercicios para el aprendizaje de la lengua náhuatl de
# Hueyapan y diccionario español-náhuatl" (CDI, 2016). Se usó para ajustar el
# borrador; solo se muestra en la hoja «Glosario», no prellena la plantilla.
GLOSARIO_HUEYAPAN = {
    "pueblo / municipio": "altepetl (pl. altepemej)", "ciudad": "ueyaltepetl", "gente, población": "tlatlaga",
    "mujer": "soatl (pl. soamej)", "hombre": "tlagatl (pl. tlagamej)", "niño": "piltontle", "joven": "telpogatl",
    "escuela": "kaltlamachtiloyan, tlamachtiloyan", "maestro": "temachtiani", "alumnos": "momachtianij",
    "enseñanza": "temachtilistle", "aprender": "salo, momachtia", "escribir": "tlajkuiloa", "leer": "amapoua",
    "casa": "kalle", "tierra": "tlale", "camino": "ojtle", "agua": "atl", "luz": "tlauille",
    "trabajo / trabajar": "teguitl / teguiti", "dinero": "tomin", "enfermedad": "kogolistle", "medicina": "pajtle",
    "contar / cuenta": "tlapoua / tlapoualli", "teléfono": "teposmecanonotsali", "platicar": "tlanonotsale",
    "ayudar / ayuda": "paleuia / tlapaleuia", "cuidar": "tlajpia", "buscar": "temo", "ver": "itta",
    "abrir / cerrar": "tlapoua / tsagua", "escoger": "pepena", "limpiar": "pojpoa", "llevar": "kuiga",
    "fiesta": "iljuitl", "nuevo": "yankuik", "grande": "ueyi", "lejos": "uejka", "hoy, ahora": "axan",
    "dónde": "kanin", "cuánto": "keski", "más": "ogachi", "algunos": "siguin", "con": "iga", "y": "uan, an",
    "no": "amo", "todo": "nochi", "vida": "nemilistle", "bien": "kuale", "silla": "ikpalli", "cama": "tlapechtle",
}
FUENTE_HUEYAPAN = "Diccionario náhuatl de Hueyapan (Montero Baeza, CDI 2016)"
FUENTE_GLOSARIO = "Glosario oficial (Traducción transformación digital.docx)"
# Traducción de apoyo generada con IA: el traductor debe revisarla.
FUENTE_BORRADOR = "Borrador automático, ajustado a Hueyapan (revisar)"

# No se traducen: nombres propios, siglas, marcas, correos y listas de campos.
NO_TRADUCIR = {
    "Emiliano Zapata", "Jojutla", "— Emiliano Zapata", "— Jojutla", "Morelos", "Cuernavaca",
    "Margarita González Saravia", "Daniel Sotelo Salazar", "ATDEM", "IEBEM", "CSV", "KML",
    "Facebook", "Instagram", "YouTube", "Digital.morelos", "atencionciudadana@morelos.gob.mx",
    "Xochitepec", "— Xochitepec", "Centro",
}
RE_LISTA_CAMPOS = re.compile(r"^[\w{}+-]+(, [\w{}+-]+)+$")   # "region, municipio, poligono…"

INSTRUCCIONES = [
    ("Plantilla de traducción al náhuatl", True),
    ("Tablero Territorios de Paz y Buen Vivir · Agencia de Transformación Digital", False),
    ("", False),
    ("1. Escriba la traducción en la columna «Traducción al náhuatl (/na/)» de la hoja «Textos».", False),
    ("2. No modifique la columna «Texto en español»: es la clave con la que el tablero encuentra cada texto.", False),
    ("3. {n} representa un número que cambia (por ejemplo «{n} localidades» = «132 localidades»). "
     "Consérvelo en la traducción, en el lugar que corresponda gramaticalmente.", False),
    ("4. {d} representa un dato que cambia (nombre de escuela, localidad, dependencia…). "
     "Consérvelo igual que {n}. Si hay varios, se llenan en el mismo orden que en español.", False),
    ("5. Los nombres de municipios, localidades, escuelas y personas no se traducen.", False),
    ("6. Si un texto no debe traducirse (sigla, nombre propio), deje la celda vacía y anótelo en «Notas».", False),
    ("7. Las filas en verde ya tienen la traducción oficial del portal morelos.gob.mx; puede corregirlas.", False),
    ("   Las filas en amarillo son un BORRADOR automático: revíselas y corríjalas; al validarlas, "
     "cambie la columna «Fuente» por «Revisada». El borrador sigue la variante de Hueyapan "
     "(Tetela del Volcán) según el diccionario de Montero Baeza (CDI, 2016); vea la hoja «Glosario».", False),
    ("8. Columna «Tipo»: «texto» es visible en pantalla; «atributo» es texto para lectores de pantalla "
     "o ayudas emergentes (no se ve, pero es importante para accesibilidad).", False),
    ("", False),
    ("Cuando termine, entregue este archivo a la ATDEM para cargarlo al tablero.", False),
]

COLUMNAS = [("N°", 6), ("Sección", 24), ("Tipo", 10), ("Texto en español", 70),
            ("Traducción al náhuatl (/na/)", 70), ("Fuente", 22), ("Notas", 30)]
COL_CLAVE, COL_TRAD, COL_FUENTE, COL_NOTAS = 4, 5, 6, 7

VERDE_OSCURO = "2E3B2B"
VERDE_CLARO = "E4E8DC"
AMARILLO_BORRADOR = "FFF4CC"


def traducible(clave: str) -> bool:
    return clave not in NO_TRADUCIR and not RE_LISTA_CAMPOS.match(clave)


def cubierto_por_patron(clave: str, patrones: list) -> bool:
    """'Ver Escuela X en el mapa' ya lo cubre 'Ver {d} en el mapa'."""
    return "{d}" not in clave and any(p.match(clave) for p in patrones)


def patron_de(clave: str):
    partes = re.split(r"(\{[nd]\})", clave)
    # En las claves los números ya son "{n}": solo {d} es comodín.
    fuente = "".join(".+?" if p == "{d}" else re.escape(p) for p in partes)
    return re.compile(f"^{fuente}$")


def leer_existentes() -> dict:
    """Traducciones y notas ya capturadas en una plantilla previa."""
    if not SALIDA.exists():
        return {}
    ws = load_workbook(SALIDA)[HOJA]
    existentes = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        clave = fila[COL_CLAVE - 1]
        if clave:
            existentes[clave] = {"trad": fila[COL_TRAD - 1], "fuente": fila[COL_FUENTE - 1], "notas": fila[COL_NOTAS - 1]}
    return existentes


def textos_a_traducir() -> list:
    textos = json.loads(ENTRADA.read_text(encoding="utf-8"))
    patrones = [patron_de(t["clave"]) for t in textos if "{d}" in t["clave"]]
    return [t for t in textos if traducible(t["clave"]) and not cubierto_por_patron(t["clave"], patrones)]


def hoja_instrucciones(wb: Workbook) -> None:
    ws = wb.create_sheet("Instrucciones", 0)
    ws.column_dimensions["A"].width = 120
    for i, (texto, titulo) in enumerate(INSTRUCCIONES, start=1):
        celda = ws.cell(row=i, column=1, value=texto)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        celda.font = Font(bold=True, size=14, color=VERDE_OSCURO) if titulo else Font(size=11)


def hoja_textos(wb: Workbook, textos: list, existentes: dict) -> dict:
    ws = wb.create_sheet(HOJA)
    encabezado = Font(bold=True, color="FFFFFF")
    for col, (nombre, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=nombre)
        celda.font = encabezado
        celda.fill = PatternFill("solid", fgColor=VERDE_OSCURO)
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[celda.column_letter].width = ancho
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:G{len(textos) + 1}"

    rellenos = {FUENTE_GLOSARIO: PatternFill("solid", fgColor=VERDE_CLARO),
                FUENTE_BORRADOR: PatternFill("solid", fgColor=AMARILLO_BORRADOR)}
    conteo = {"conservadas": 0, "glosario": 0}
    for n, t in enumerate(textos, start=1):
        clave = t["clave"]
        previa = existentes.get(clave, {})
        trad, fuente = previa.get("trad"), previa.get("fuente")
        if trad:
            conteo["conservadas"] += 1
        elif clave in GLOSARIO:
            trad, fuente = GLOSARIO[clave], FUENTE_GLOSARIO
            conteo["glosario"] += 1
        tipo = "texto" if t["tipo"] == "texto" else "atributo"
        fila = [n, t["seccion"], tipo, clave, trad, fuente, previa.get("notas")]
        for col, valor in enumerate(fila, start=1):
            celda = ws.cell(row=n + 1, column=col, value=valor)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            relleno = rellenos.get(FUENTE_BORRADOR if (fuente or "").startswith("Borrador") else fuente)
            if relleno and col == COL_TRAD:
                celda.fill = relleno
    return conteo


def hoja_glosario(wb: Workbook) -> None:
    ws = wb.create_sheet("Glosario")
    ws.append(["Español", "Náhuatl (/na/)"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE_OSCURO)
    ws.cell(row=1, column=3, value="Fuente").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=3).fill = PatternFill("solid", fgColor=VERDE_OSCURO)
    for es, na in GLOSARIO.items():
        ws.append([es, na, "Glosario oficial del portal"])
    for es, na in GLOSARIO_HUEYAPAN.items():
        ws.append([es, na, FUENTE_HUEYAPAN])
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60


def main() -> None:
    if not ENTRADA.exists():
        raise SystemExit(f"No existe {ENTRADA}. Ejecute antes scripts/extraer_textos_nahuatl.js")
    existentes = leer_existentes()
    textos = textos_a_traducir()
    wb = Workbook()
    wb.remove(wb.active)
    hoja_instrucciones(wb)
    conteo = hoja_textos(wb, textos, existentes)
    hoja_glosario(wb)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    palabras = sum(len(t["clave"].split()) for t in textos)
    perdidas = [k for k, v in existentes.items() if v.get("trad") and k not in {t["clave"] for t in textos}]
    print(f"Plantilla: {SALIDA.relative_to(RAIZ)}")
    print(f"  {len(textos)} textos ({palabras} palabras); {conteo['glosario']} prellenados del glosario; "
          f"{conteo['conservadas']} traducciones conservadas")
    if perdidas:
        print(f"  AVISO: {len(perdidas)} traducciones previas ya no aparecen en el tablero: {perdidas[:5]}")


if __name__ == "__main__":
    main()
